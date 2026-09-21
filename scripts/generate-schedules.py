"""Converte a planilha de horários da escola em JSON para a área admin.

Uso: python scripts/generate-schedules.py origem.xlsx data/horarios-turmas.json
"""

import json
import re
import sys
from collections import Counter
from datetime import time
from pathlib import Path

from openpyxl import load_workbook


CLASS_NAMES = {
    "6º ANO A": "6º A", "6º ANO B": "6º B", "6º ANO C": "6º C",
    "7º ANO A": "7º A", "7º ANO B": "7º B", "7º ANO C": "7º C",
    "7º ANO D": "7º D", "7º ANO E": "7º E", "7º ANO F": "7º F",
    "8º ANO A": "8º A", "8º ANO B": "8º B", "8º ANO C": "8º C",
    "8º ANO D": "8º D", "8º ANO E": "8º E",
    "9º ANO A": "9º A", "9º ANO B": "9º B", "9º ANO C": "9º C",
    "9º ANO D": "9º D", "9º ANO E": "9º E",
    "1ª SÉRIE A BENS DE CONSUMO": "1º Neg A",
    "1ª SÉRIE B BENS DE CONSUMO": "1º Neg B",
    "1ª SÉRIE C TECH": "1º Tec C", "1ª SÉRIE D TECH": "1º Tec D",
    "1ª SÉRIE E TECH": "1º Tec E", "1ª SÉRIE F TECH": "1º Tec F",
    "1ª SÉRIE G TECH": "1º Tec G",
    "2ª SÉRIE A FINTECH": "2º Fin A",
    "2ª SÉRIE B BENS DE CONSUMO": "2º Neg B",
    "2ª SÉRIE C BENS DE CONSUMO": "2º Neg C",
    "2ª SÉRIE D TECH": "2º Tec D", "2ª SÉRIE E TECH": "2º Tec E",
    "2ª SÉRIE F TECH": "2º Tec F", "2ª SÉRIE G TECH": "2º Tec G",
    "2ª SÉRIE H TECH": "2º Tec H", "2ª SÉRIE I TECH": "2º Tec I",
    "3ª SÉRIA A FINTECH": "3º Fin A", "3ª SÉRIE B FINTECH": "3º Fin B",
    "3ª SÉRIE C BENS DE CONSUMO": "3º Neg C",
    "3ª SÉRIE D BENS DE CONSUMO": "3º Neg D",
    "3ª SÉRIE E TECH DE": "3º Tec DS E",
    "3ª SÉRIE F TECH DE": "3º Tec DS F",
    "3ª SÉRIE G TECH DE": "3º Tec DS G",
    "3ª SÉRIE H TECH AD": "3º Tec AD H",
    "3ª SÉRIE I TECH AD": "3º Tec AD I",
}

def clean(value):
    return " ".join(str(value or "").replace("\u00a0", " ").split())


def title_key(value):
    return clean(value).upper()


def main(source, destination):
    workbook = load_workbook(source, read_only=True, data_only=True)
    records = []
    skipped = Counter()
    unmapped_headers = set()
    seen = set()
    discovered_classes = set()

    for sheet in workbook.worksheets:
        current_class = None
        for row in range(1, sheet.max_row + 1):
            heading = title_key(sheet.cell(row, 4).value)
            if re.search(r"(?:ANO|SÉRIE|SÉRIA) [A-I](?: |$)", heading):
                current_class = CLASS_NAMES.get(heading)
                if current_class:
                    discovered_classes.add(current_class)
                else:
                    unmapped_headers.add(heading)
                continue

            period = clean(sheet.cell(row, 1).value)
            start = sheet.cell(row, 2).value
            end = sheet.cell(row, 3).value
            if not (re.fullmatch(r"\d+[ªº]", period) and isinstance(start, time) and isinstance(end, time)):
                continue
            if current_class is None:
                skipped["fora das turmas escolhidas"] += 1
                continue

            for day in range(1, 6):
                subject = clean(sheet.cell(row, day + 3).value)
                professor = clean(sheet.cell(row + 1, day + 3).value)
                if not subject or subject in {"---", "-"}:
                    skipped["sem matéria"] += 1
                    continue
                if not professor or professor in {"---", "-"} or professor.startswith("(B)"):
                    skipped["sem professor"] += 1
                    continue
                record = {
                    "professor": professor,
                    "subject": subject,
                    "className": current_class,
                    "floor": "",
                    "startTime": start.strftime("%H:%M"),
                    "endTime": end.strftime("%H:%M"),
                    "roomDescription": "",
                    "dayOfWeek": day,
                    "active": True,
                }
                fingerprint = tuple(record[field] for field in ("className", "dayOfWeek", "startTime", "endTime", "professor", "subject"))
                if fingerprint not in seen:
                    records.append(record)
                    seen.add(fingerprint)
                else:
                    skipped["duplicado exato"] += 1

    if unmapped_headers:
        raise ValueError(f"Turmas sem mapeamento: {sorted(unmapped_headers)}")
    expected_classes = set(CLASS_NAMES.values())
    if discovered_classes != expected_classes:
        raise ValueError(f"Turmas ausentes na planilha: {sorted(expected_classes - discovered_classes)}")

    records.sort(key=lambda item: (item["className"], item["dayOfWeek"], item["startTime"], item["professor"]))
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"records": len(records), "classes": len(discovered_classes), "skipped": skipped}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Uso: python scripts/generate-schedules.py origem.xlsx destino.json")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
